import pyodbc
import openpyxl
from openpyxl.styles import Border, Side, Font, Alignment
from openpyxl.utils import get_column_letter
import configparser
import os
import sys
from datetime import datetime

def get_base_path():
    """Возвращает путь к папке, где находится исполняемый файл (или скрипт)"""
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    else:
        return os.path.dirname(os.path.abspath(__file__))

def read_config(config_file='config.ini'):
    """Читает параметры подключения из config.ini (рядом с exe)"""
    base = get_base_path()
    config_path = os.path.join(base, config_file)
    config = configparser.ConfigParser()
    config.read(config_path, encoding='utf-8')
    if 'DB' not in config:
        raise ValueError("В config.ini отсутствует секция [DB]")
    return config['DB']

def read_query(query_file='query.sql'):
    """Читает SQL-запрос из файла (рядом с exe)"""
    base = get_base_path()
    query_path = os.path.join(base, query_file)
    if not os.path.exists(query_path):
        raise FileNotFoundError(f"Файл {query_path} не найден")
    with open(query_path, 'r', encoding='utf-8') as f:
        return f.read()

def build_connection_string(config):
    """Формирует строку подключения на основе конфига"""
    driver = config.get('driver', 'ODBC Driver 17 for SQL Server')
    server = config.get('server')
    database = config.get('database')
    if not server or not database:
        raise ValueError("В конфиге должны быть указаны server и database")
    conn_str = f'DRIVER={{{driver}}};SERVER={server};DATABASE={database};'
    trusted = config.get('trusted_connection', '').lower() in ('yes', 'true', '1')
    if trusted:
        conn_str += 'Trusted_Connection=yes;'
    else:
        username = config.get('username')
        password = config.get('password')
        if username and password:
            conn_str += f'UID={username};PWD={password};'
        else:
            raise ValueError("Не указаны данные для аутентификации. Используйте trusted_connection или username/password.")
    return conn_str

def fetch_data(conn_str, query):
    """Выполняет запрос, возвращает (колонки, список строк)"""
    conn = pyodbc.connect(conn_str, timeout=30)
    cursor = conn.cursor()
    cursor.execute(query)
    # Получаем имена колонок
    columns = [desc[0] for desc in cursor.description] if cursor.description else []
    # Все строки
    rows = cursor.fetchall()
    conn.close()
    return columns, rows

def save_to_excel(columns, rows, output_file):
    """Сохраняет данные в Excel с форматированием (границы, автоширина, заголовки)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'

    # Заголовки
    if columns:
        for col_idx, col_name in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = Font(bold=True)

    # Данные
    if rows:
        for row_idx, row in enumerate(rows, 2):
            for col_idx, value in enumerate(row, 1):
                # openpyxl умеет работать с datetime, None, числами, строками
                ws.cell(row=row_idx, column=col_idx, value=value)

    # Стили (границы, выравнивание)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    center_alignment = Alignment(horizontal='center', vertical='center')

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = thin_border
            cell.alignment = center_alignment
            if cell.row == 1:
                cell.font = Font(bold=True)  # повторно для заголовков

    # Автоширина колонок
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                try:
                    max_len = max(max_len, len(str(cell.value)))
                except:
                    pass
        ws.column_dimensions[col_letter].width = max_len + 3

    wb.save(output_file)

def main():
    try:
        print("Чтение конфигурации...")
        config = read_config()
        conn_str = build_connection_string(config)

        print("Чтение SQL-запроса...")
        query = read_query()
        if not query.strip():
            raise ValueError("Файл запроса пуст")

        print("Выполнение запроса к БД...")
        columns, rows = fetch_data(conn_str, query)

        if not rows:
            print("Предупреждение: запрос не вернул данных. Будет создан пустой файл с заголовками (если есть).")

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        output_file = f"result_{timestamp}.xlsx"

        print("Сохранение в Excel...")
        save_to_excel(columns, rows, output_file)

        print(f"Готово! Результат сохранён в {output_file}")
        print(f"Количество записей: {len(rows)}")

    except FileNotFoundError as e:
        print(f"Ошибка: {e}")
        sys.exit(1)
    except ValueError as e:
        print(f"Ошибка в конфигурации: {e}")
        sys.exit(1)
    except pyodbc.Error as e:
        print(f"Ошибка подключения к БД: {e}")
        sys.exit(1)
    except Exception as e:
        print(f"Непредвиденная ошибка: {e}")
        sys.exit(1)

if __name__ == '__main__':
    main()